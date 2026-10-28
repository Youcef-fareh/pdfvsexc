"""
Batch Runner — processes a queue of FileEntry objects sequentially.
Emits progress via a callback so the UI can update without blocking.
"""
from __future__ import annotations

import os
import traceback
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum, auto
from typing import Callable, Optional

from app.pipeline.input_handler import FileEntry, InputType
from app.utils.config import EXPORTS_DIR


class BatchStatus(Enum):
    PENDING = auto()
    PROCESSING = auto()
    SUCCESS = auto()
    FAILED = auto()
    SKIPPED = auto()


@dataclass
class BatchResult:
    file_entry: FileEntry
    status: BatchStatus = BatchStatus.PENDING
    output_path: str = ""
    error_message: str = ""
    processing_seconds: float = 0.0
    avg_confidence: float = 0.0
    page_count: int = 0


ProgressCallback = Callable[[int, int, BatchResult], None]
# signature: (current_index, total_count, result_so_far) → None


# ── Runner ────────────────────────────────────────────────────────────────────

def run_batch(
    entries: list[FileEntry],
    export_format: str = "xlsx",
    export_mode: str = "clone",
    template_name: str | None = None,
    output_dir: str = EXPORTS_DIR,
    progress_cb: ProgressCallback | None = None,
    cancel_flag: list[bool] | None = None,   # cancel_flag[0] = True to abort
) -> list[BatchResult]:
    """
    Process each FileEntry sequentially and return a list of BatchResults.

    cancel_flag: pass a mutable list [False] — set [0] = True from UI to cancel.
    """
    results: list[BatchResult] = []
    total = len(entries)

    for idx, entry in enumerate(entries):
        if cancel_flag and cancel_flag[0]:
            remaining = BatchResult(file_entry=entry, status=BatchStatus.SKIPPED)
            results.append(remaining)
            if progress_cb:
                progress_cb(idx + 1, total, remaining)
            continue

        result = BatchResult(file_entry=entry, status=BatchStatus.PROCESSING)
        start = datetime.now()

        try:
            result = _process_one(
                entry=entry,
                export_format=export_format,
                export_mode=export_mode,
                template_name=template_name,
                output_dir=output_dir,
            )
        except Exception as exc:
            result.status = BatchStatus.FAILED
            result.error_message = f"{type(exc).__name__}: {exc}"

        result.processing_seconds = (datetime.now() - start).total_seconds()
        results.append(result)

        if progress_cb:
            progress_cb(idx + 1, total, result)

    return results


# ── Single-file processing ────────────────────────────────────────────────────

def _process_one(
    entry: FileEntry,
    export_format: str,
    export_mode: str,
    template_name: str | None,
    output_dir: str,
) -> BatchResult:
    """Run the full pipeline for a single file and write the output."""
    # Lazy imports so startup is fast
    from app.pipeline.preprocessor import preprocess
    from app.pipeline.ocr_engine import ocr_page
    from app.pipeline.table_detector import detect_table
    from app.pipeline.normalizer import normalize_table
    from app.utils.pdf_utils import pdf_to_images, extract_tables_from_pdf
    from app.export.xlsx_writer import write_xlsx
    from app.export.docx_writer import write_docx

    result = BatchResult(file_entry=entry)
    all_grids: list = []
    confidences: list[float] = []

    if entry.input_type == InputType.DIGITAL_PDF:
        # Direct extraction — no OCR needed
        tables = extract_tables_from_pdf(entry.path, password=entry.password)
        all_grids.extend(tables)
        confidences.append(100.0)

    else:
        # Image or scanned PDF — run through OCR pipeline
        if entry.input_type == InputType.IMAGE:
            from PIL import Image as PILImage
            images = [(1, PILImage.open(entry.path))]
        else:
            images = list(pdf_to_images(entry.path, password=entry.password))

        for page_num, pil_img in images:
            preprocessed = preprocess(pil_img)
            ocr_result = ocr_page(preprocessed)
            confidences.append(ocr_result.average_confidence)
            grid = detect_table(preprocessed, ocr_result)
            if grid:
                all_grids.extend(grid)

    if not all_grids:
        result.status = BatchStatus.FAILED
        result.error_message = "No table data could be extracted from this file."
        return result

    # Flatten / normalize
    from app.pipeline.normalizer import NormalizedTable, normalize_table
    combined_grid = all_grids if isinstance(all_grids[0], list) else [all_grids]
    norm_table = normalize_table(combined_grid)

    # Export
    base_name = os.path.splitext(os.path.basename(entry.path))[0]
    os.makedirs(output_dir, exist_ok=True)

    if export_format == "xlsx":
        out_path = os.path.join(output_dir, f"{base_name}.xlsx")
        write_xlsx(norm_table, out_path, mode=export_mode)
    else:
        out_path = os.path.join(output_dir, f"{base_name}.docx")
        write_docx(norm_table, out_path, mode=export_mode)

    result.status = BatchStatus.SUCCESS
    result.output_path = out_path
    result.avg_confidence = sum(confidences) / len(confidences) if confidences else 0.0
    result.page_count = entry.page_count
    return result


# ── Batch report ──────────────────────────────────────────────────────────────

def generate_batch_report(results: list[BatchResult], output_dir: str) -> str:
    """Write a summary XLSX report of all batch results. Returns file path."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Batch Report"

    headers = ["File", "Status", "Pages", "Avg Confidence %", "Time (s)", "Output", "Error"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A1A2E")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    status_colors = {
        BatchStatus.SUCCESS: "C8E6C9",
        BatchStatus.FAILED: "FFCDD2",
        BatchStatus.SKIPPED: "FFF9C4",
    }

    for row_idx, r in enumerate(results, 2):
        color = status_colors.get(r.status, "FFFFFF")
        fill = PatternFill("solid", fgColor=color)
        row_data = [
            os.path.basename(r.file_entry.path),
            r.status.name,
            r.page_count,
            f"{r.avg_confidence:.1f}",
            f"{r.processing_seconds:.2f}",
            r.output_path,
            r.error_message,
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = fill

    # Auto-width
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].auto_size = True

    report_path = os.path.join(output_dir, "batch_report.xlsx")
    wb.save(report_path)
    return report_path
